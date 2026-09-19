#!/usr/bin/env python3
"""
Build the speaker-notes round-trip document from the deck sources.

    python3 tools/build_speaker_notes.py          # writes .md and .docx, then self-tests
    python3 tools/build_speaker_notes.py --check  # self-test an existing .md only

The .docx needs `python-docx`. Without it you still get the .md.

WHY YOU WANT THIS. Writing speaker notes inside a .qmd is unpleasant, and a
coauthor or an advisor cannot review them there at all. This pulls every
`::: {.notes}` block into one Markdown file and one Word file, in deck order.
Edit that document, hand it back, and `apply_speaker_notes.py` writes the
changes into the .qmd files.

The round trip only works if the document parses back EXACTLY, so the last thing
this script does is parse its own output and compare it to the source, character
for character. If that self-test ever fails, fix it before trusting the applier.

THE HEADING LINE IS THE CONTRACT:

    ## [12] #method - How the study was run

The id is what an applier keys on. Everything from one heading to the next is
that slide's block. Position and title are there for a human.

SOURCE OF TRUTH IS THE .qmd, not the rendered HTML, because the notes are edited
back into `::: {.notes}` blocks. Deck order comes from expanding `{{< include >}}`
in place, which is how Quarto builds it. The title slide is the exception: Quarto
generates it from the YAML, so its notes live in title-slide-notes.html and are
lifted from that file's JS array.
"""
import re, os, sys, html as _html, json, datetime

# tools/ lives one level below the project root.
ROOT   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DECK   = 'talk'
MASTER = os.path.join(ROOT, 'master')
STAMP  = os.environ.get('NOTES_DATE') or datetime.date.today().isoformat()
MD     = os.path.join(ROOT, 'speaker_notes_%s.md' % STAMP)
DOCX   = os.path.join(ROOT, 'speaker_notes_%s.docx' % STAMP)

# A divider slide has no heading text, so name it here by its id. Add your own.
LABELS = {
    'study1': 'Divider — Study 1', 'study2': 'Divider — Study 2',
    'discussion': 'Divider — Discussion',
    'appendix': 'Divider — Appendix',
    'thanks': 'Thank you / Questions',
}


def expand(path, seen=None):
    """The .qmd with includes expanded in place, i.e. in deck order."""
    seen = seen or set()
    out = []
    for line in open(path):
        m = re.match(r'\s*\{\{<\s*include\s+(\S+)\s*>\}\}', line)
        if m:
            inc = os.path.join(os.path.dirname(path), m.group(1))
            if inc not in seen:
                seen.add(inc)
                out.append(expand(inc, seen))
            continue
        out.append(line)
    return ''.join(out)


def parse_slides():
    src = expand(os.path.join(MASTER, DECK + '.qmd'))
    src = re.sub(r'\A---\n.*?\n---\n', '', src, flags=re.S)      # YAML header
    lines = src.split('\n')
    slides, i, fence = [], 0, False
    while i < len(lines):
        ln = lines[i]
        if ln.startswith('```'):
            fence = not fence; i += 1; continue
        if not fence and ln.startswith('## '):
            head = ln[3:].strip()
            m = re.match(r'^(.*?)\s*\{(.*)\}\s*$', head)
            title, attrs = (m.group(1).strip(), m.group(2)) if m else (head, '')
            sid = re.search(r'#([A-Za-z0-9_-]+)', attrs)
            sid = sid.group(1) if sid else ''
            uncounted = 'visibility="uncounted"' in attrs
            j, notes, body, depth, collecting, f2 = i + 1, [], [], 0, False, False
            while j < len(lines):
                l2 = lines[j]
                if l2.startswith('```'):
                    f2 = not f2
                    if collecting: notes.append(l2)
                    else: body.append(l2)
                    j += 1; continue
                if not f2 and l2.startswith('## '):
                    break
                if not f2 and re.match(r'^:::+\s*\{\.notes\}\s*$', l2):
                    collecting, depth = True, 1; j += 1; continue
                if collecting:
                    if not f2 and re.match(r'^:::+\s*\{', l2):
                        depth += 1
                    elif not f2 and re.match(r'^:::+\s*$', l2):
                        depth -= 1
                        if depth == 0:
                            collecting = False; j += 1; continue
                    notes.append(l2)
                else:
                    body.append(l2)
                j += 1
            slides.append({'title': title, 'id': sid, 'uncounted': uncounted,
                           'notes': '\n'.join(notes).strip(),
                           'body': '\n'.join(body).strip()})
            i = j; continue
        i += 1

    # Title slide: notes live in the JS array in title-slide-notes.html.
    tsn = ''
    p = os.path.join(ROOT, 'includes', 'title-slide-notes.html')
    if os.path.exists(p):
        txt = open(p).read()
        m = re.search(r'aside\.innerHTML\s*=\s*\[(.*?)\]\.join\(', txt, flags=re.S)
        if m:
            # Re-join the fragments the way the page does. Splitting on quotes
            # without rejoining truncates sentences that span array elements.
            frags = re.findall(r"'((?:[^'\\]|\\.)*)'", m.group(1), flags=re.S)
            tsn = ''.join(f.replace("\\'", "'") for f in frags)
        tsn = re.sub(r'<li>', '- ', tsn)
        tsn = re.sub(r'</li>', '\n', tsn)
        tsn = re.sub(r'</?(ol|ul)>', '\n', tsn)
        tsn = re.sub(r'<code>(.*?)</code>', r'`\1`', tsn, flags=re.S)
        tsn = re.sub(r'</p>\s*', '\n\n', tsn)
        tsn = re.sub(r'<br\s*/?>', '\n', tsn)
        tsn = re.sub(r'<strong>(.*?)</strong>', r'**\1**', tsn, flags=re.S)
        tsn = re.sub(r'<em>(.*?)</em>', r'*\1*', tsn, flags=re.S)
        tsn = _html.unescape(re.sub(r'<[^>]+>', '', tsn)).strip()
    slides.insert(0, {'title': '(generated title slide)',
                      'id': 'title-slide', 'uncounted': False, 'notes': tsn,
                      'body': 'Generated by Quarto from the YAML header.'})

    shown = 0
    for n, s in enumerate(slides, start=1):
        s['pos'] = n
        if not s['uncounted']:
            shown += 1; s['shown'] = shown
        else:
            s['shown'] = None
    return slides


def onscreen(body):
    """What the audience actually reads. Markup and authoring notes removed."""
    b = body or ''
    b = re.sub(r'<!--.*?-->', '', b, flags=re.S)
    b = re.sub(r'^```\{r.*?^```', '[figure]', b, flags=re.S | re.M)
    b = re.sub(r'^```.*$', '', b, flags=re.M)
    b = re.sub(r'<style.*?</style>', '', b, flags=re.S | re.I)
    b = re.sub(r'<script.*?</script>', '', b, flags=re.S | re.I)
    b = re.sub(r'\{\{<\s*include.*?>\}\}', '[included figure]', b)
    b = re.sub(r'!\[\]\((.*?)\)\{[^}]*\}', r'[image: \1]', b)
    b = re.sub(r'^:::.*$', '', b, flags=re.M)
    b = re.sub(r'<[^>]+>', ' ', b)
    b = _html.unescape(b)
    b = re.sub(r'\{[^}\n]*\}', '', b)
    out = []
    for line in b.split('\n'):
        t = ' '.join(line.split())
        if not t:
            if out and out[-1] != '': out.append('')
            continue
        if len(re.sub(r'[^A-Za-z0-9]', '', t)) < 2:
            continue
        out.append(t)
    return '\n'.join(out).strip()


HOWTO = """HOW TO USE IT
1. Edit the notes freely. Rewrite, cut, add, reorder within a slide.
2. Do NOT change a slide's heading line. That line is how the notes get matched
   back to the right slide. Everything else in a slide's block is yours.
3. Hand the whole file back and it goes into the deck in one pass.
4. To add notes to a slide that has none, just write them under its heading.

WHAT THE HEADING MEANS
  [12] #fade - What you set up early is what still shows
   12    = position in the deck, counting the title slide as 1.
   #fade = the slide's id, which is what the deck actually keys on.

Appendix slides are marked UNCOUNTED. They sit outside the audience-facing slide
count and are only reached by jumping to them during questions.

THE 'ON SCREEN' BLOCK IS REFERENCE ONLY. It shows what the audience sees, so you
can edit notes without the deck open. Changes there are NOT applied. If you want
a slide's visible text changed, say so separately."""


def label(s):
    t = s['title'].replace('&nbsp;', '').strip()
    return t or LABELS.get(s['id'], '(untitled)')


def write_md(slides):
    ncount = sum(1 for s in slides if not s['uncounted'])
    md = ['# Speaker notes', '',
          '%d slides, %d counted plus %d appendix.'
          % (len(slides), ncount, len(slides) - ncount),
          '', '---', '', '```', HOWTO, '```', '', '---', '']
    for s in slides:
        tag = 'UNCOUNTED' if s['uncounted'] else 'slide %d of %d' % (s['shown'], ncount)
        md += ['## [%d] #%s - %s' % (s['pos'], s['id'], label(s)), '', '*%s*' % tag, '']
        scr = onscreen(s['body'])
        if scr:
            md.append('> **ON SCREEN (reference only, not applied):**')
            md += ['> ' + l if l.strip() else '>' for l in scr.split('\n')]
            md.append('')
        md += ['**NOTES**', '',
               s['notes'] if s['notes'] else '_(no speaker notes yet — write them here)_',
               '', '---', '']
    open(MD, 'w').write('\n'.join(md))
    return MD


def parse_back(path):
    """The applier's parse. Kept here so the self-test exercises the real thing."""
    md = open(path).read()
    blocks = re.split(r'^## \[(\d+)\] #([A-Za-z0-9_-]+) - (.*)$', md, flags=re.M)
    out = {}
    for i in range(1, len(blocks), 4):
        sid, body = blocks[i + 1], blocks[i + 3]
        idx = body.find('**NOTES**')
        notes = body[idx + len('**NOTES**'):] if idx >= 0 else ''
        # The trailing `---` is the document's own separator. It has no leading
        # newline when the notes were emptied, which is how an emptied block used
        # to parse back as the literal string '---'.
        notes = re.sub(r'(?:\n|\A)---\s*$', '', notes.strip()).strip()
        if notes.startswith('_(no speaker notes'):
            notes = ''
        out[sid] = notes
    return out


def selftest(slides, path):
    """Parse our own output back and demand it match the source exactly."""
    src = {s['id']: s['notes'].strip() for s in slides}
    got = parse_back(path)
    missing = set(src) - set(got)
    extra = set(got) - set(src)
    bad = [sid for sid in src if got.get(sid, '').strip() != src[sid]]
    ok = not missing and not extra and not bad
    print('  slides: %d in source, %d parsed back' % (len(src), len(got)))
    print('  chars : %d in, %d out' % (sum(len(v) for v in src.values()),
                                       sum(len(v) for v in got.values())))
    if missing: print('  MISSING:', sorted(missing))
    if extra:   print('  EXTRA  :', sorted(extra))
    if bad:     print('  MISMATCHED:', bad[:10])
    print('  ROUND TRIP:', 'OK' if ok else 'FAILED')
    return ok


def write_docx(slides):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
    except ImportError:
        print('  python-docx not installed; skipped the .docx')
        return None
    GREEN, GOLD, GREY = RGBColor(0x1a,0x56,0x32), RGBColor(0x7A,0x5A,0x02), RGBColor(0x66,0x66,0x66)
    ncount = sum(1 for s in slides if not s['uncounted'])
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.add_heading('Speaker notes', 0)
    p = doc.add_paragraph()
    p.add_run('%d slides: %d counted, %d appendix. Generated from the .qmd sources on %s.'
              % (len(slides), ncount, len(slides) - ncount, STAMP))
    doc.add_heading('How to use this', level=1)
    for line in HOWTO.split('\n'):
        if line.strip():
            doc.add_paragraph(line)

    def emit(par, text):
        for tok in re.split(r'(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`)', text):
            if not tok: continue
            if tok.startswith('**') and tok.endswith('**'): par.add_run(tok[2:-2]).bold = True
            elif tok.startswith('*') and tok.endswith('*'): par.add_run(tok[1:-1]).italic = True
            elif tok.startswith('`') and tok.endswith('`'):
                r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'; r.font.size = Pt(9.5)
            else: par.add_run(tok)

    for s in slides:
        doc.add_page_break()
        hp = doc.add_heading('', level=1)
        hp.add_run('[%d]  #%s  —  %s' % (s['pos'], s['id'], label(s))).font.color.rgb = GREEN
        tag = 'UNCOUNTED (appendix — jump target only)' if s['uncounted'] \
              else 'slide %d of %d' % (s['shown'], ncount)
        p = doc.add_paragraph(); r = p.add_run(tag)
        r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GREY
        scr = onscreen(s['body'])
        if scr:
            p = doc.add_paragraph(); r = p.add_run('ON SCREEN — reference only, not applied')
            r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = GOLD
            for line in scr.split('\n'):
                if not line.strip(): continue
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(line.strip()); r.font.size = Pt(8.5); r.font.color.rgb = GREY
        p = doc.add_paragraph(); r = p.add_run('NOTES'); r.bold = True; r.font.color.rgb = GREEN
        if s['notes']:
            for block in s['notes'].split('\n\n'):
                if block.strip():
                    emit(doc.add_paragraph(), block.strip().replace('\n', ' '))
        else:
            r = doc.add_paragraph().add_run('(no speaker notes yet — write them here)')
            r.italic = True; r.font.color.rgb = GREY
    doc.save(DOCX)
    return DOCX


if __name__ == '__main__':
    slides = parse_slides()
    ncount = sum(1 for s in slides if not s['uncounted'])
    print('parsed %d slides (%d counted, %d appendix)' % (len(slides), ncount, len(slides) - ncount))
    empty = [s['id'] for s in slides if not s['notes']]
    print('slides with no notes:', empty or 'none')
    if '--check' in sys.argv:
        sys.exit(0 if selftest(slides, MD) else 1)
    print('wrote', write_md(slides))
    d = write_docx(slides)
    if d: print('wrote', d)
    print('self-test:')
    sys.exit(0 if selftest(slides, MD) else 1)
